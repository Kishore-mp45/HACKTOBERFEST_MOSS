from pathlib import Path
from zipfile import BadZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.errors import FileNotDecryptedError, PdfReadError

from moss.parsers.docx_parser import DocxParser
from moss.parsers.markdown_parser import MarkdownParser
from moss.parsers.pdf_parser import PdfParser
from moss.parsers.text_parser import TextParser


def test_text_parser_handles_utf8_and_latin1(tmp_path: Path) -> None:
    utf8 = tmp_path / "utf8.txt"
    utf8.write_text("Moss handles café notes.", encoding="utf-8")
    latin1 = tmp_path / "latin1.txt"
    latin1.write_bytes("Résumé".encode("latin-1"))

    assert TextParser().parse(utf8).text == "Moss handles café notes."
    assert TextParser().parse(latin1).text == "Résumé"


def test_text_parser_handles_empty_files(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")

    result = TextParser().parse(path)

    assert result.text == ""
    assert result.word_count == 0
    assert result.char_count == 0


def test_markdown_parser_strips_yaml_front_matter(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text(
        "---\ntitle: private metadata\n---\n# Project notes\nBody text",
        encoding="utf-8",
    )

    result = MarkdownParser().parse(path)

    assert "private metadata" not in result.text
    assert result.text.startswith("# Project notes")


def test_markdown_parser_preserves_document_without_front_matter(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("--- this is a heading\nbody", encoding="utf-8")

    assert MarkdownParser().parse(path).text.startswith("--- this is a heading")


def test_markdown_parser_handles_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "empty.md"
    path.write_text("", encoding="utf-8")

    assert MarkdownParser().parse(path).text == ""


def test_docx_parser_preserves_paragraphs_and_metadata(tmp_path: Path) -> None:
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.core_properties.title = "Planning notes"
    document.save(path)

    result = DocxParser().parse(path)

    assert result.text == "First paragraph\nSecond paragraph"
    assert result.metadata["title"] == "Planning notes"


def test_docx_parser_handles_empty_document(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    Document().save(path)

    result = DocxParser().parse(path)

    assert result.text == ""
    assert result.word_count == 0


def test_docx_parser_rejects_corrupted_document(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a DOCX archive")

    with pytest.raises(BadZipFile):
        DocxParser().parse(path)


def test_docx_parser_rejects_truncated_document(tmp_path: Path) -> None:
    path = tmp_path / "truncated.docx"
    path.write_bytes(b"PK\x03\x04\x14\x00")

    with pytest.raises(BadZipFile):
        DocxParser().parse(path)


def test_pdf_parser_returns_page_text_chunks(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as output:
        writer.write(output)

    result = PdfParser().parse(path)

    assert result.page_count == 2
    assert result.page_chunks == [{"page": 1, "text": ""}, {"page": 2, "text": ""}]


def test_corrupted_pdf_raises_a_file_specific_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a PDF")

    with pytest.raises(PdfReadError):
        PdfParser().parse(path)


@pytest.mark.parametrize("suffix", [".txt", ".pdf", ".docx"])
def test_parsers_handle_zero_byte_inputs(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"empty{suffix}"
    path.write_bytes(b"")
    parser = {
        ".txt": TextParser(),
        ".pdf": PdfParser(),
        ".docx": DocxParser(),
    }[suffix]

    if suffix == ".txt":
        assert parser.parse(path).text == ""
    elif suffix == ".pdf":
        with pytest.raises(PdfReadError):
            parser.parse(path)
    else:
        with pytest.raises(BadZipFile):
            parser.parse(path)


def test_pdf_parser_rejects_encrypted_document_without_password(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("password")
    with path.open("wb") as output:
        writer.write(output)

    with pytest.raises(FileNotDecryptedError):
        PdfParser().parse(path)


def test_pdf_page_errors_do_not_expose_error_details(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import moss.parsers.pdf_parser as pdf_parser

    class BrokenPage:
        def extract_text(self) -> str:
            raise ValueError("private document excerpt")

    class Reader:
        pages = [BrokenPage()]
        metadata = None

        def __init__(self, _path: Path, strict: bool = False) -> None:
            assert strict is False

    monkeypatch.setattr(pdf_parser, "PdfReader", Reader)

    result = PdfParser().parse(tmp_path / "ignored.pdf")

    assert result.error == "Page 1: ValueError"
    assert "private document excerpt" not in (result.error or "")
