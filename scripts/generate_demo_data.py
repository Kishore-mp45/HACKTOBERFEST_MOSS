"""Create deterministic, synthetic demo workspaces for local demos and submissions.

Supports:
- Rich multi-format workspace (default): 18 documents matching the Moss UI demo fixtures
  (PDFs, Word DOCX, Excel XLSX, Markdown, text files, byte-identical duplicates, and scanned docs).
- Flat benchmark workspace (--flat): 30 synthetic Markdown documents across three topics.
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path
from typing import TypedDict

import docx  # type: ignore[import-untyped]
import openpyxl  # type: ignore[import-untyped]
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


class ProjectContent(TypedDict):
    topic: str
    details: list[str]
    actions: list[str | None]


FLAT_PROJECTS: dict[str, ProjectContent] = {
    "cedar": {
        "topic": "Cedar product launch",
        "details": [
            "The design team reviewed the updated onboarding flow.",
            "The launch checklist tracks decisions from the weekly planning meeting.",
            "The team is comparing two options for the first-run experience.",
            "Release notes should explain the changes in plain language.",
            "Feedback from the pilot group is being collected for review.",
            "The support guide needs a short section about account setup.",
            "The product group agreed to keep the initial release focused.",
            "A follow-up meeting will review the pilot results.",
            "The team is organizing research notes before the launch.",
            "The onboarding work is coordinated across design and support.",
        ],
        "actions": [
            "TODO: Send the revised onboarding draft to the design team.",
            "Deadline: Review the launch checklist by November 6, 2026.",
            "TODO: Summarize pilot feedback for the product meeting.",
            None,
            "TODO: Prepare a short support guide for account setup.",
            None,
            "Deadline: Confirm the release notes by November 13, 2026.",
            None,
            "TODO: Share the research summary with the team.",
            None,
        ],
    },
    "harbor": {
        "topic": "Harbor community garden",
        "details": [
            "The north beds are being prepared for the winter planting plan.",
            "The volunteer group discussed water access and tool storage.",
            "The garden map records the current paths and shared beds.",
            "The compost area needs clearer signs for new volunteers.",
            "Seasonal notes describe which plants did well this year.",
            "The group is planning a small seed exchange for neighbors.",
            "The tool shed inventory was reviewed by the coordinators.",
            "The garden calendar includes community work sessions.",
            "The path between the beds should remain accessible.",
            "Volunteer notes are collected after each work session.",
        ],
        "actions": [
            "TODO: Ask volunteers to label the shared tools.",
            None,
            "Deadline: Update the garden map by November 8, 2026.",
            "TODO: Print clear signs for the compost area.",
            None,
            "TODO: Prepare a list of seeds for the exchange.",
            None,
            "Deadline: Publish the winter work calendar by November 15, 2026.",
            None,
            "TODO: Collect volunteer feedback after the next session.",
        ],
    },
    "orchard": {
        "topic": "Orchard finance planning",
        "details": [
            "The quarterly budget compares planned and actual expenses.",
            "The finance notes group invoices by project and month.",
            "The forecast includes expected maintenance costs.",
            "The team reviewed the process for recording reimbursements.",
            "Monthly totals are checked before the report is shared.",
            "The budget worksheet keeps a record of approved estimates.",
            "Finance meeting notes include questions about next quarter.",
            "The invoice archive is organized by vendor and date.",
            "The project forecast was updated with current estimates.",
            "Expense summaries are prepared for the planning group.",
        ],
        "actions": [
            "Deadline: Reconcile October invoices by November 10, 2026.",
            None,
            "TODO: Confirm the maintenance estimate with the facilities team.",
            None,
            "TODO: Review reimbursement entries before month end.",
            None,
            "Deadline: Share the next-quarter forecast by November 18, 2026.",
            None,
            "TODO: File the updated project estimate.",
            None,
        ],
    },
}


def make_pdf(
    path: Path,
    title: str,
    pages_lines: list[list[str]],
    *,
    has_text_layer: bool = True,
) -> None:
    """Generate a clean, valid PDF file."""
    path.parent.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()

    for lines in pages_lines:
        page = writer.add_blank_page(width=612, height=792)
        if not has_text_layer:
            continue

        stream = DecodedStreamObject()
        content = ["BT", "/F1 11 Tf", "15 TL", "54 720 Td"]
        safe_title = title.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content.append(f"({safe_title}) Tj")
        content.append("T*")
        content.append("T*")

        for line in lines:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            content.append(f"({safe}) '")

        content.append("ET")
        stream.set_data("\n".join(content).encode("latin1"))

        resources = DictionaryObject()
        fonts = DictionaryObject()
        font = DictionaryObject()
        font[NameObject("/Type")] = NameObject("/Font")
        font[NameObject("/Subtype")] = NameObject("/Type1")
        font[NameObject("/BaseFont")] = NameObject("/Helvetica")
        fonts[NameObject("/F1")] = font
        resources[NameObject("/Font")] = fonts

        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = stream

    with open(path, "wb") as f:
        writer.write(f)


def make_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    """Generate a clean Word DOCX document."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        if p.startswith("## "):
            doc.add_heading(p[3:], level=2)
        elif p.startswith("### "):
            doc.add_heading(p[4:], level=3)
        else:
            doc.add_paragraph(p)
    doc.save(str(path))


def make_xlsx(
    path: Path,
    title: str,
    headers: list[str],
    rows: list[list[str | float | int]],
) -> None:
    """Generate a clean Excel XLSX workbook."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def make_text(path: Path, content: str) -> None:
    """Write text or markdown with UTF-8 encoding."""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.strip() + "\n", encoding="utf-8", newline="\n")


def generate_rich_demo_workspace(output_directory: Path) -> list[Path]:
    """Generate the 18 rich multi-format files that match the Moss UI fixtures."""
    output_directory.mkdir(parents=True, exist_ok=True)
    created: list[Path] = []

    # 1. Research/climate-paper-v3.pdf
    f1 = output_directory / "Research" / "climate-paper-v3.pdf"
    make_pdf(
        f1,
        "Urban Heat Island Dynamics Across 14 European Metropolitan Areas",
        [
            [
                "Authors: Dr. Elena Vance, Dr. Julian Meier, Institute for Urban Climatology",
                "Preprint v3 - Published: October 2026",
                "",
                "Abstract: This empirical study investigates the thermal gradient of urban heat islands",
                "across 14 major European cities. Using satellite radiometric sensing and ground sensors,",
                "we analyze heat-island effects across 14 European cities, with an open dataset and mitigation scenarios.",
                "",
                "1. Introduction",
                "Metropolitan centers experience microclimate shifts characterized by nighttime heat retention.",
                "Surface permeability and canopy coverage represent the strongest mitigation factors.",
            ],
            [
                "3. Methodology & Normalisation",
                "3.2 Methodology: adapted from the open EU dataset (see climate-paper-v3).",
                "Data collection protocols were harmonized across all 14 sensor grids.",
                "Temperature deltas were normalized relative to rural reference baselines.",
                "",
                "4. Key Findings",
                "The findings confirm high diurnal variance in dense commercial cores.",
                "Urban green corridors reduce localized air temperature by up to 2.4 degrees Celsius.",
                "The accompanying dataset is released under Open Data Commons license for academic reuse.",
            ],
        ],
    )
    created.append(f1)

    # 2. Downloads/climate-paper-v3 (1).pdf (Exact duplicate)
    f2 = output_directory / "Downloads" / "climate-paper-v3 (1).pdf"
    f2.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(f1, f2)
    created.append(f2)

    # 3. Finance/annual-report-2024-final.pdf
    f3 = output_directory / "Finance" / "annual-report-2024-final.pdf"
    make_pdf(
        f3,
        "Annual Financial & Operational Report - Fiscal Year 2024",
        [
            [
                "Executive Summary & Consolidated Financial Statements",
                "Compiled annual report: revenue by segment, cost structure, and auditor notes.",
                "Audited in accordance with International Financial Reporting Standards.",
                "",
                "Revenue Analysis:",
                "- Enterprise Software: 4.82M EUR (+18% YoY)",
                "- Consulting & Maintenance Services: 2.14M EUR",
                "- Cloud Infrastructure & Hosting: 1.05M EUR",
                "",
                "Q3 expense totals reconcile with the accompanying spreadsheet.",
            ],
            [
                "Notes to the Financial Statements",
                "Cost structure highlights capital expenditure reduction in on-premise compute.",
                "Auditor notes confirm all liabilities were settled within standard payment windows.",
                "Q3 totals carried into FY summary (see annual report, cost structure).",
                "",
                "Independent Auditor Confirmation: Clean unqualified opinion issued.",
            ],
        ],
    )
    created.append(f3)

    # 4. Finance/annual-report-2024-final (copy).pdf (Exact duplicate)
    f4 = output_directory / "Finance" / "annual-report-2024-final (copy).pdf"
    f4.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(f3, f4)
    created.append(f4)

    # 5. Thesis/thesis-ch3-urban-heat-draft.docx
    f5 = output_directory / "Thesis" / "thesis-ch3-urban-heat-draft.docx"
    make_docx(
        f5,
        "Chapter 3: Urban Heat Island Mitigation Strategies in High-Density Districts",
        [
            "Doctoral Dissertation - Draft Working Copy",
            "## 3.1 Introduction and Theoretical Framework",
            "Urban microclimates are heavily shaped by geometry and surface albedo. This chapter investigates district-level cooling interventions.",
            "## 3.2 Methodology: adapted from the open EU dataset (see climate-paper-v3).",
            "The comparative analysis builds directly on the EU urban climate study dataset across 14 cities. Normalisation follows standard rural baseline offsets.",
            "## 3.3 Normalisation and Metric Definitions",
            "Thermal flux metrics are normalized per hectare of impervious surface cover.",
            "## 3.4 Spatial Analysis Results",
            "TODO: complete regression modeling for section 3.4 and finalize temperature delta tables before the supervisor review meeting.",
            "## 3.5 Discussion and Policy Implications",
            "Zoning requirements incorporating green canopy corridors provide sustainable cooling.",
        ],
    )
    created.append(f5)

    # 6. Thesis/thesis-notes-ideas.md
    f6 = output_directory / "Thesis" / "thesis-notes-ideas.md"
    make_text(
        f6,
        """# Thesis Working Notes & Ideas

Working notes for the thesis: outline for chapter 3, open questions about dataset normalisation, and reminders to email the supervisor.

## Chapter 3 Progress & Outline
- Ch.3 outline — reuse the normalisation approach from the EU heat-island study
- Compare empirical sensor measurements from Berlin and Vienna
- Check dataset normalisation for seasonal variations

## Questions for Supervisor
- Clarify whether to include winter thermal sensor data in Section 3.2
- TODO: email supervisor regarding methodology review
""",
    )
    created.append(f6)

    # 7. Work/meeting-notes-2025-06-12.md
    f7 = output_directory / "Work" / "meeting-notes-2025-06-12.md"
    make_text(
        f7,
        """# Team Sync & Planning Notes — 12 June 2025

Team sync notes: budget revision for Q3, hiring update, and personal follow-ups captured at the end.

## Agenda Items
1. Q3 Budget Reconciliation
2. Infrastructure & Hosting Migration
3. Acme GmbH Vendor Contract Status

## Meeting Notes
- Discussed updated operational costs across cloud infrastructure.
- Reviewed Acme GmbH frame contract; services end in September 2026.
- The revised Q3 budget figures are ready for team sign-off.

## Action Items
- TODO: send revised Q3 budget to Anja before Monday
- Personal: finally book that dentist appointment
""",
    )
    created.append(f7)

    # 8. Apartment/lease-agreement-signed.pdf
    f8 = output_directory / "Apartment" / "lease-agreement-signed.pdf"
    make_pdf(
        f8,
        "Residential Lease Agreement - Signed",
        [
            [
                "Residential Tenancy Agreement",
                "Landlord: Hausverwaltung Mitte GmbH",
                "Tenant: Authorized Resident",
                "Property: Bergmannstrasse 12, 10961 Berlin",
                "",
                "Clause 4: Operational Expenses and Maintenance",
                "Clause 4.2: Maintenance responsibilities for common areas and minor heating fixtures",
                "shall be shared equally between landlord and tenant. Heating repairs are reimbursable.",
                "",
                "Notice Period: 3 months to the end of any calendar month.",
                "Signed and dated in Berlin.",
            ]
        ],
    )
    created.append(f8)

    # 9. Apartment/rent-invoice-2026-06.pdf
    f9 = output_directory / "Apartment" / "rent-invoice-2026-06.pdf"
    make_pdf(
        f9,
        "Rent & Utility Statement - June 2026",
        [
            [
                "Invoice Number: INV-2026-06-882",
                "Statement Period: 01.06.2026 - 30.06.2026",
                "Address: Bergmannstrasse 12, 10961 Berlin",
                "",
                "Line Items:",
                "- Base Monthly Rent: 980.00 EUR",
                "- Heating & Warm Water Advance: 140.00 EUR",
                "- Operating Costs Advance: 75.00 EUR",
                "",
                "Maintenance costs are shared per clause 4.2 of the rental contract.",
                "Total Amount Due: 1,195.00 EUR",
            ]
        ],
    )
    created.append(f9)

    # 10. Apartment/insurance-policy-home-2026.pdf
    f10 = output_directory / "Apartment" / "insurance-policy-home-2026.pdf"
    make_pdf(
        f10,
        "Comprehensive Household Insurance Certificate 2026",
        [
            [
                "Policy Reference: POL-DE-2026-8941",
                "Insurer: Berlin Household & Fire Assurance AG",
                "Insured location: Bergmannstrasse 12, 10961 Berlin.",
                "",
                "Coverage Specifications:",
                "- Water damage and pipe burst: Included up to 100,000 EUR",
                "- Storm and elemental hazards: Included",
                "- Glass breakage: Included",
                "",
                "Coverage runs until the renewal window in November 2026.",
            ]
        ],
    )
    created.append(f10)

    # 11. Travel/flight-itinerary-BER-LIS.pdf
    f11 = output_directory / "Travel" / "flight-itinerary-BER-LIS.pdf"
    make_pdf(
        f11,
        "Flight Booking Confirmation: Berlin (BER) - Lisbon (LIS)",
        [
            [
                "Airline: European Air Mobility",
                "Booking Reference: LH-88392",
                "Passenger: Resident Traveler",
                "",
                "Outbound Flight: BER -> LIS, 2 November 2026",
                "Return: LIS -> BER, 9 November.",
                "",
                "Status: Confirmed. Seat selection and online check-in window opens 24h prior.",
            ]
        ],
    )
    created.append(f11)

    # 12. Travel/lisbon-stay-confirmation.pdf
    f12 = output_directory / "Travel" / "lisbon-stay-confirmation.pdf"
    make_pdf(
        f12,
        "Lisbon Accommodation Confirmation",
        [
            [
                "Reservation ID: LIS-HOST-4491",
                "Property: Santa Catarina Rooftop Suite, Lisbon",
                "Stay: 2-9 November, host contact inside.",
                "",
                "Check-in Window: 15:00 - 20:00 on 2 November",
                "Host Contact: Miguel Santos (+351 912 345 678)",
                "Address: Rua de Santa Catarina 45, 1200 Lisbon",
                "",
                "Wi-Fi and access code provided upon arrival verification.",
            ]
        ],
    )
    created.append(f12)

    # 13. Downloads/scan_2026-03-14_1132.pdf (Scanned PDF without text layer)
    f13 = output_directory / "Downloads" / "scan_2026-03-14_1132.pdf"
    make_pdf(
        f13,
        "Scanned Image Document",
        [[]],
        has_text_layer=False,
    )
    created.append(f13)

    # 14. Work/q3-expense-report-DRAFT-final2.xlsx
    f14 = output_directory / "Work" / "q3-expense-report-DRAFT-final2.xlsx"
    make_xlsx(
        f14,
        "Q3 Expenses",
        ["Date", "Category", "Vendor", "Description", "Amount EUR", "Approved", "Notes"],
        [
            ["2026-07-15", "Travel", "Lufthansa", "Flight booking BER-LIS", 240.50, "Yes", "Travel budget"],
            ["2026-08-02", "Hardware", "Dell Computers", "Server SSD replacements", 850.00, "Yes", "Infrastructure"],
            ["2026-08-18", "Services", "Acme GmbH", "Maintenance services", 1200.00, "Yes", "Acme contract"],
            ["2026-09-10", "Software", "CloudFlare", "DNS and security annual renewal", 480.00, "Yes", "Approved"],
            ["2026-09-30", "Summary", "Finance Group", "Q3 totals carried into FY summary (see annual report, cost structure).", 2770.50, "Yes", "Reconciled"],
        ],
    )
    created.append(f14)

    # 15. Work/vendor-contract-acme.docx
    f15 = output_directory / "Work" / "vendor-contract-acme.docx"
    make_docx(
        f15,
        "Master Services Agreement: Acme GmbH",
        [
            "Frame contract with Acme GmbH for maintenance services.",
            "Term ended 30 September 2026; no renewal clause was triggered.",
            "## 1. Scope of Work",
            "Acme GmbH provided maintenance and operational support for on-premise hardware.",
            "## 2. Invoicing and Payments",
            "Invoices were billed on a net-30 schedule. All pending vouchers were cleared.",
            "## 3. Expiration and Archival",
            "The agreement concluded on 30 September 2026. This contract is ready for historical archiving.",
        ],
    )
    created.append(f15)

    # 16. Misc/downloads-todo-backlog.txt
    f16 = output_directory / "Misc" / "downloads-todo-backlog.txt"
    make_text(
        f16,
        """Personal Backlog & Errand Checklist:
- renew library card
- compare bike insurance quotes
- file the warranty papers for the espresso machine
- verify cloud backup drives
""",
    )
    created.append(f16)

    # 17. Misc/photo-backup-inventory.md
    f17 = output_directory / "Misc" / "photo-backup-inventory.md"
    make_text(
        f17,
        """# Photo Backup & Archive Inventory

Inventory of photo backups across two drives and one cloud bucket, with folder sizes and last-verified dates.

- Drive 1 (Seagate 4TB External): Last verified August 2026, 1.4 TB photos and raw negatives.
- Drive 2 (Western Digital 2TB Mirror): Stored at home office.
- Cloud Bucket (Backblaze B2): Encrypted monthly cold storage mirror.

TODO: verify cloud bucket restore before December
""",
    )
    created.append(f17)

    # 18. Thesis/old/thesis-ch3-draft-old.docx
    f18 = output_directory / "Thesis" / "old" / "thesis-ch3-draft-old.docx"
    make_docx(
        f18,
        "Chapter 3: Urban Heat Island Mitigation (Archived Early Draft)",
        [
            "Early working draft - Superseded by current chapter draft.",
            "## 3.1 Initial Notes",
            "Preliminary thoughts on temperature sensors and dataset structures.",
            "Contains outdated formulas superseded by the finalized methodology in thesis-ch3-urban-heat-draft.docx.",
        ],
    )
    created.append(f18)

    return created


def document_text(project: str, topic: str, detail: str, action: str | None) -> str:
    lines = [
        f"# {topic}",
        "",
        f"Synthetic demo note for the {project} workspace collection.",
        "",
        detail,
        "",
    ]
    if action:
        lines.extend([action, ""])
    lines.extend(
        [
            "This fabricated content is provided only for local demonstration.",
            "It contains no real people, organizations, or commitments.",
            "",
        ]
    )
    return "\n".join(lines)


def generate_flat_demo_data(output_directory: Path) -> list[Path]:
    output_directory.mkdir(parents=True, exist_ok=True)
    generated: list[Path] = []
    for project, content in FLAT_PROJECTS.items():
        for index, (detail, action) in enumerate(
            zip(content["details"], content["actions"], strict=True),
            start=1,
        ):
            path = output_directory / f"{project}-{index:02d}.md"
            expected = document_text(project, content["topic"], detail, action)
            path.write_text(expected, encoding="utf-8", newline="\n")
            generated.append(path)
    return generated


def generate_demo_data(
    output_directory: Path,
    *,
    mode: str = "rich",
) -> list[Path]:
    """Generate demo workspace files in either rich or flat mode."""
    if mode == "flat":
        return generate_flat_demo_data(output_directory)
    return generate_rich_demo_workspace(output_directory)


def main() -> None:
    repository_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=repository_root / "demo_workspace",
        help="Output directory (default: ./demo_workspace)",
    )
    parser.add_argument(
        "--flat",
        action="store_true",
        help="Generate 30 flat markdown documents instead of the rich multi-format workspace",
    )
    args = parser.parse_args()
    mode = "flat" if args.flat else "rich"
    target = args.output.expanduser().resolve()
    generated = generate_demo_data(target, mode=mode)
    print(f"Generated: {len(generated)} {mode} synthetic demo documents in {target}")


if __name__ == "__main__":
    main()
